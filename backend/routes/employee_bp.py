import os, uuid, pandas as pd
from werkzeug.utils import secure_filename
from io import BytesIO
from datetime import datetime, timedelta
from flask import Blueprint, request, jsonify, send_file
from sqlalchemy import or_, func, and_
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation


from extensions import db

from model.blacklist import OsBlacklist
from model.employment import OsEmployment
from model.person import OsPerson
from model.subCompany import SubCompany
from model.costCenter import costCenter
from model.canteen import canteen, canteenDetail
from model.osType import osType
from model.card import OsCard
from model.osCostCenter import OsCostCenter
from model.grade import OsGrade
from model.alokasi import Alokasi
from model.ob_emp import ObEmployee
from model.hr_models import User, UserSubcompanyAccess

employee_bp = Blueprint('employee_bp', __name__)

UPLOAD_FOLDER = 'static/uploads/photos'
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

@employee_bp.teardown_request
def teardown_request(exception=None):
    try:
        db.session.remove()
    except Exception:
        pass

def clean_str(val):
    if val is None or pd.isna(val):
        return None
    return str(val).strip()

def parse_use_cc(val):
    """Utility untuk konversi input flag use_cc menjadi integer (0 atau 1)"""
    if str(val).strip() in ['1', 'true', 'True']:
        return 1
    return 0

def parse_date(date_str):
    """Konversi string tanggal menjadi objek Date, kembalikan None jika kosong/tidak valid (Cegah MySQL Error 1292)"""
    if not date_str:
        return None
    s = str(date_str).strip()
    if s.lower() in ('', 'null', 'none', 'undefined', 'nan', 'nat'):
        return None
    try:
        return datetime.strptime(s, '%Y-%m-%d').date()
    except ValueError:
        return None

def extract_excel_date(val):
    """Proteksi ekstra untuk membaca format tanggal dari Excel/Pandas"""
    if pd.isna(val):
        return None
    if hasattr(val, 'date'):
        return val.date()
    return parse_date(str(val))

def get_allowed_subcompanies():
    """Helper internal untuk mengambil daftar subcompany_id yang diizinkan untuk user aktif"""
    user_email = request.headers.get('X-User-Email')
    if not user_email:
        return []
    user = User.query.filter_by(email=user_email).first()
    if not user:
        return []
    access_records = UserSubcompanyAccess.query.filter_by(user_id=user.id).all()
    return [a.sub_company_id for a in access_records]

def process_and_save_photo(file_storage, target_folder, filename_without_ext, max_width=600, quality=80):
    try:
        img = Image.open(file_storage)
        try:
            img = ImageOps.exif_transpose(img)
        except Exception:
            pass
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        if img.width > max_width:
            ratio = max_width / float(img.width)
            new_height = int((float(img.height) * float(ratio)))
            img = img.resize((max_width, new_height), Image.Resampling.LANCZOS)
        new_filename = f"{filename_without_ext}.webp"
        file_path = os.path.join(target_folder, new_filename)
        img.save(file_path, "WEBP", quality=quality, optimize=True)
        return f"/{target_folder}/{new_filename}"
    except Exception as e:
        print(f"[ERROR] Gagal mengompresi foto: {str(e)}")
        return None

@employee_bp.route('/employee')
def index():
    try:
        page = request.args.get('page', 1, type=int)
        pageSize = request.args.get('pageSize', 10, type=int)
        search = request.args.get('search', '', type=str)

        status = request.args.get('status', 'all', type=str)
        sub_company_id = request.args.get('sub_company', '', type=str)
        department_id = request.args.get('department', '', type=str)
        target_date_str = request.args.get('target_date', '', type=str)

        # 1. Tentukan Titik Waktu (Point-in-Time)
        if target_date_str:
            target_date = datetime.strptime(target_date_str, '%Y-%m-%d').date()
        else:
            target_date = datetime.now().date()

        # 2. OPTIMASI ENTERPRISE: POINT-IN-TIME EXPLICIT JOIN (Anti N+1 Query)
        query = db.session.query(
            OsEmployment,
            OsPerson.name.label('person_name'),
            costCenter.org_name.label('cc_name'),
            OsCostCenter.cc_id.label('cc_code'),
            OsCard.card_number,
            OsGrade.grade,
            osType.type_worker,
            osType.posisi
        ).join(
            OsPerson, OsEmployment.person_id == OsPerson.person_id
        ).outerjoin(
            OsCostCenter,
            and_(
                OsCostCenter.employee_id == OsEmployment.id,
                OsCostCenter.valid_from <= target_date,
                or_(OsCostCenter.valid_to >= target_date, OsCostCenter.valid_to == None)
            )
        ).outerjoin(
            costCenter, costCenter.id == OsCostCenter.org_cc_id
        ).outerjoin(
            OsCard,
            and_(
                OsCard.employee_id == OsEmployment.id,
                OsCard.valid_from <= target_date,
                or_(OsCard.valid_to >= target_date, OsCard.valid_to == None)
            )
        ).outerjoin(
            OsGrade,
            and_(
                OsGrade.employee_id == OsEmployment.id,
                OsGrade.valid_from <= target_date,
                or_(OsGrade.valid_to >= target_date, OsGrade.valid_to == None)
            )
        ).outerjoin(
            osType,
            and_(
                osType.employee_id == OsEmployment.id,
                osType.valid_from <= target_date,
                or_(osType.valid_to >= target_date, osType.valid_to == None)
            )
        )

        # 3. LOGIKA FILTER HAK AKSES SUBCOMPANY (SSO)
        allowed_subcos = get_allowed_subcompanies()
        if allowed_subcos:
            query = query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))
            if sub_company_id and sub_company_id not in allowed_subcos and sub_company_id not in ['TYPE_OS', 'TYPE_VENDOR']:
                return jsonify({"status": "error", "message": "Akses ditolak"}), 403

        # 4. Filter Pencarian
        if search:
            query = query.filter(
                or_(
                    OsEmployment.employee_code.cast(db.String).ilike(f"%{search}%"),
                    OsPerson.name.ilike(f"%{search}%"),
                    OsCard.card_number.ilike(f"%{search}%")
                )
            )

        # 5. LOGIKA FILTER AKTIF BERDASARKAN POINT-IN-TIME
        
        if status == 'active':
            query = query.filter(
                and_(
                    OsEmployment.valid_from.is_not(None),
                    OsEmployment.valid_from <= target_date,
                    or_(OsEmployment.valid_to >= target_date, OsEmployment.valid_to == None)
                )
            )
        elif status == 'inactive':
            query = query.filter(
                or_(
                    OsEmployment.valid_from.is_(None),
                    OsEmployment.valid_to < target_date
                )
            )

        # 6. FILTER DINAMIS
        if sub_company_id == 'TYPE_OS':
            subquery_os = db.session.query(SubCompany.sub_company_id).filter(SubCompany.type_company == 'OS')
            query = query.filter(OsEmployment.sub_company_id.in_(subquery_os))            
        elif sub_company_id == 'TYPE_VENDOR':
            subquery_vendor = db.session.query(SubCompany.sub_company_id).filter(SubCompany.type_company == 'Vendor')
            query = query.filter(OsEmployment.sub_company_id.in_(subquery_vendor))            
        elif sub_company_id:
            query = query.filter(OsEmployment.sub_company_id == sub_company_id)

        if department_id:
            query = query.filter(OsCostCenter.org_cc_id == department_id)
        
        # 7. Eksekusi Pagination
        pagination = query.paginate(page=page, per_page=pageSize, error_out=False)

        # 8. MAPPING DATA
        result_data = []
        for emp, person_name, cc_name, cc_code, card_number, grade, type_worker, posisi in pagination.items:
            emp_dict = emp.to_dict() 
            
            # Override data fluktuatif (SCD Type 2) dari Explicit JOIN Point-in-Time
            emp_dict['person_name'] = person_name
            emp_dict['cc_name'] = cc_name if cc_name else '-'
            emp_dict['cost_center_id'] = cc_code if cc_code else '-'
            emp_dict['card_number'] = card_number if card_number else '-'
            emp_dict['grade'] = grade if grade else '-'
            emp_dict['type_worker'] = type_worker if type_worker else '-'
            emp_dict['posisi'] = posisi if posisi else '-'
            
            result_data.append(emp_dict)

        return jsonify({
            "status": "success",
            "data": result_data,
            "total_page": pagination.pages,
            "current_page": pagination.page,
            "total_item": pagination.total
        }), 200
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500


@employee_bp.route('/employee/search-autocomplete', methods=['GET'])
def search_autocomplete():
    try:
        query_str = request.args.get('q', '').strip()
        if len(query_str) < 3:
            return jsonify({"status": "success", "data": []})
        
        today = datetime.now().date()
        base_query = db.session.query(OsEmployment, OsPerson)\
            .join(OsPerson, OsEmployment.person_id == OsPerson.person_id)

        allowed_subcos = get_allowed_subcompanies()
        if allowed_subcos:
            base_query = base_query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))

        results = base_query.filter(
            or_(
                OsPerson.name.ilike(f"%{query_str}%"),
                OsEmployment.employee_code.cast(db.String).ilike(f"%{query_str}%")
            )
        ).limit(20).all()
        
        data_result = []
        for emp, person in results:
            is_active = bool(emp.valid_from and emp.valid_from <= today and (emp.valid_to is None or emp.valid_to >= today))
            
            data_result.append({
                "emp_pk_id": emp.id,
                "employee_code": emp.employee_code,
                "name": person.name,
                "resident_id": person.resident_id,
                "use_cc": getattr(emp, 'use_cc', 0),
                "is_active": is_active,
                "status_text": "Aktif" if is_active else "Non-Aktif"
            })
        
        return jsonify({"status": "success", "data": data_result}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@employee_bp.route('/employee/search-all', methods=['GET'])
def search_all():
    try:
        query_str = request.args.get('q', '').strip()
        if len(query_str) < 3:
            return jsonify({"status": "success", "data": []})        
        today = datetime.now().date()
        data_result = []

        base_os_query = db.session.query(OsEmployment, OsPerson)\
            .join(OsPerson, OsEmployment.person_id == OsPerson.person_id)

        allowed_subcos = get_allowed_subcompanies()
        if allowed_subcos:
            base_os_query = base_os_query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))

        results_os = base_os_query.filter(
            or_(
                OsPerson.name.ilike(f"%{query_str}%"),
                OsEmployment.employee_code.cast(db.String).ilike(f"%{query_str}%")
            )
        ).all()

        for emp, person in results_os:
            is_active = bool(emp.valid_from and emp.valid_from <= today and (emp.valid_to is None or emp.valid_to >= today))

            data_result.append({
                "source": "OS",
                "emp_pk_id": emp.id,
                "employee_code": emp.employee_code,
                "name": person.name,
                "use_cc": getattr(emp, 'use_cc', 0),
                "is_active": is_active,
                "status_text": "Aktif" if is_active else "Non-Aktif"
            })

        if not allowed_subcos:
            results_ob = ObEmployee.query.filter(
                or_(
                    ObEmployee.employee_name.ilike(f"%{query_str}%"),
                    ObEmployee.employee_id.cast(db.String).ilike(f"%{query_str}%")
                )
            ).all()

            for ob in results_ob:
                data_result.append({
                    "source": "OB",
                    "emp_pk_id": ob.employee_id,
                    "employee_code": ob.employee_id,
                    "name": ob.employee_name,
                    "use_cc": 0,
                    "is_active": True,
                    "status_text": "Aktif"
                })
        
        return jsonify({"status": "success", "data": data_result}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@employee_bp.route('/employee/search/<string:emp_id>', methods=['GET'])
def search_employee(emp_id):
    try:
        query = db.session.query(OsPerson.name, OsEmployment.id) \
            .join(OsEmployment, OsPerson.person_id == OsEmployment.person_id) \
            .filter(OsEmployment.employee_code == emp_id)
        
        allowed_subcos = get_allowed_subcompanies()
        if allowed_subcos:
            query = query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))

        result = query.first()
        if result:
            return jsonify({"status": "success", "full_name": result.name, "emp_pk_id": result.id}), 200
        return jsonify({"status": "error", "message": "Employee ID tidak ditemukan atau akses ditolak"}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@employee_bp.route('/employee/submit', methods=['POST'])
def add():
    try:
        data = request.json if request.is_json else request.form

        if not data.get('nama') or str(data.get('nama')).strip() == "":
            return jsonify({"status": "error", "message": "Nama wajib diisi!"}), 400

        if not data.get('employee_id') or str(data.get('employee_id')).strip() == "":
            return jsonify({"status": "error", "message": "ID Karyawan wajib diisi!"}), 400

        new_start_date = parse_date(data.get('valid_from'))
        if not new_start_date:
             return jsonify({"status": "error", "message": "Tanggal Valid From wajib diisi dan format harus benar!"}), 400             
        
        new_valid_to = parse_date(data.get('valid_to'))
        c_valid_from = parse_date(data.get('c_valid_from'))
        c_valid_to = parse_date(data.get('c_valid_to'))
        
        adjusted_valid_to = new_start_date - timedelta(days=1)
        employee_code_input = clean_str(data.get('employee_id'))
        card_number_input = clean_str(data.get('card_number'))
        use_cc_input = parse_use_cc(data.get('use_cc', 0))

        if card_number_input and card_number_input.lower() != 'none':
            duplicate_card = OsCard.query.filter(
                OsCard.card_number == card_number_input,
                (OsCard.valid_to >= new_start_date) | (OsCard.valid_to == None)
            ).first()
            if duplicate_card:
                raise Exception(f"Kartu nomor {card_number_input} sudah aktif digunakan oleh record lain.")
        
        person_id = data.get('person_id')
        if not person_id or person_id == "" or str(person_id).lower() == "undefined":
            target_person = OsPerson(
                name = data.get('nama'),
                gender = data.get('gender'),            
                pob = data.get('pob'),
                dob = parse_date(data.get('dob')),
                religion = data.get('religion'),
                resident_id = data.get('resident_id'),
                address = data.get('address')
            )
        else:
            target_person = OsPerson.query.get(person_id)
            target_person.name = data.get('nama', target_person.name)
            target_person.gender = data.get('gender', target_person.gender)
            target_person.pob = data.get('pob', target_person.pob)
            target_person.dob = parse_date(data.get('dob')) or target_person.dob
            target_person.religion = data.get('religion', target_person.religion)
            target_person.resident_id = data.get('resident_id', target_person.resident_id)

        if 'photo' in request.files:
            file = request.files['photo']
            if file.filename != '':
                upload_date = datetime.now().strftime('%Y%m%d')
                base_name = f"{employee_code_input}_{upload_date}"                
                saved_path = process_and_save_photo(
                    file_storage=file,
                    target_folder=UPLOAD_FOLDER,
                    filename_without_ext=base_name,
                    max_width=600,
                    quality=80
                )
                if saved_path:
                    target_person.photo = saved_path

        db.session.add(target_person)
        db.session.flush()
        person_id = target_person.person_id        

        check_nrp_owner = OsEmployment.query.filter(
            OsEmployment.employee_code == employee_code_input,
            (OsEmployment.valid_to >= new_start_date) | (OsEmployment.valid_to == None)
        ).first()        
        if check_nrp_owner and check_nrp_owner.person_id != person_id:
            raise Exception(f"NRP / ID Karyawan '{employee_code_input}' sudah terdaftar milik orang lain.")

        active_emp_for_person = OsEmployment.query.filter(
            OsEmployment.person_id == person_id,
            (OsEmployment.valid_to >= new_start_date) | (OsEmployment.valid_to == None)
        ).first()
        if active_emp_for_person and active_emp_for_person.employee_code != employee_code_input:
            raise Exception(f"Karyawan '{target_person.name}' masih berstatus AKTIF dengan NRP lama ({active_emp_for_person.employee_code}).")

        existing_active_emp = OsEmployment.query.filter(
            OsEmployment.employee_code == employee_code_input,
            (OsEmployment.valid_to >= new_start_date) | (OsEmployment.valid_to == None)
        ).first()

        if existing_active_emp and existing_active_emp.valid_from == new_start_date:
            raise Exception(f"Karyawan dengan ID {employee_code_input} sudah memiliki data aktif di tanggal yang sama.")

        if existing_active_emp:
            active_emp_pk_id = existing_active_emp.id
            target_models = [
                {"model": OsEmployment, "field": "person_id",   "val": person_id},
                {"model": OsCard,       "field": "employee_id", "val": active_emp_pk_id},
                {"model": OsGrade,      "field": "employee_id", "val": active_emp_pk_id},
                {"model": osType,       "field": "employee_id", "val": active_emp_pk_id},
                {"model": OsCostCenter, "field": "employee_id", "val": active_emp_pk_id},
                {"model": Alokasi,      "field": "employee_id", "val": active_emp_pk_id},
            ]

            for item in target_models:                
                Model = item["model"]
                field_name = item["field"]
                value = item["val"]

                if value is None:
                    continue

                old_records = Model.query.filter(
                    getattr(Model, field_name) == value,
                    (Model.valid_to >= new_start_date) | (Model.valid_to == None),
                    Model.valid_from <= adjusted_valid_to
                ).all()

                for rec in old_records:
                    rec.valid_to = adjusted_valid_to
                    db.session.add(rec)
            
        db.session.flush()

        newEmployment = OsEmployment(
            employee_code = employee_code_input,
            sub_company_id = data.get('sub_company_id'),
            person_id = person_id,
            use_cc = use_cc_input,
            valid_from = new_start_date,
            valid_to = new_valid_to
        )
        db.session.add(newEmployment)
        db.session.flush()

        if card_number_input and card_number_input.lower() != 'none':
            newCard = OsCard(
                employee_id = newEmployment.id,
                card_number = card_number_input,
                valid_from = c_valid_from,
                valid_to = c_valid_to
            )
            db.session.add(newCard)

        newGrade = OsGrade(
            employee_id = newEmployment.id,
            grade = clean_str(data.get('grade')),
            valid_from = new_start_date,
            valid_to = new_valid_to
        )        
        db.session.add(newGrade)

        newType = osType(
            employee_id = newEmployment.id,
            type_worker = clean_str(data.get('type_worker')),
            posisi = clean_str(data.get('posisi')),
            valid_from = new_start_date,
            valid_to = new_valid_to
        )
        db.session.add(newType)

        selected_cc_id = data.get('cc_id')
        if selected_cc_id:
            master_cc = costCenter.query.get(selected_cc_id)
            if master_cc:
                newCC = OsCostCenter(
                    employee_id = newEmployment.id,
                    cc_id = master_cc.cost_center,
                    org_cc_id = master_cc.id,
                    valid_from = new_start_date,
                    valid_to = new_valid_to
                )        
                db.session.add(newCC)

                cc_def = canteen.query.join(canteenDetail, canteen.canteen_id == canteenDetail.canteen_id)\
                                      .filter(canteenDetail.org_cc_id == master_cc.id).first()
                if cc_def:
                    newAlokasi = Alokasi(
                        employee_id = newEmployment.id,
                        canteen_id = cc_def.canteen_id,
                        valid_from = new_start_date,
                        valid_to = new_valid_to
                    )
                    db.session.add(newAlokasi)        
        
        db.session.commit()
        return jsonify({"status": "success", "message": "Data berhasil disimpan!"}), 201      

    except Exception as e:
        db.session.rollback()
        return jsonify({"status": "error", "message": "Terjadi kesalahan pada server: " + str(e)}), 500


@employee_bp.route('/employee/<int:id>', methods=['PUT'])
def edit(id):
    try:
        data = request.json if request.is_json else request.form

        if not data.get('nama') or str(data.get('nama')).strip() == "":
            return jsonify({"status": "error", "message": "Nama wajib diisi!"}), 400

        if not data.get('employee_id') or str(data.get('employee_id')).strip() == "":
            return jsonify({"status": "error", "message": "ID Karyawan wajib diisi!"}), 400

        employee_code_input = clean_str(data.get('employee_id'))
        card_number_input = clean_str(data.get('card_number'))
        use_cc_input = parse_use_cc(data.get('use_cc', 0))
        
        new_start_date = parse_date(data.get('valid_from'))
        new_valid_to = parse_date(data.get('valid_to'))
        c_valid_from = parse_date(data.get('c_valid_from'))
        c_valid_to = parse_date(data.get('c_valid_to'))        
        day_before = new_start_date - timedelta(days=1) if new_start_date else None

        target_emp = OsEmployment.query.get(id)
        if not target_emp:
            return jsonify({"status": "error", "message": "Data Employment tidak ditemukan!"}), 404

        if card_number_input and card_number_input.lower() != 'none' and new_start_date:
            duplicate_card = OsCard.query.filter(
                OsCard.card_number == card_number_input,
                OsCard.employee_id != id,
                ((OsCard.valid_to >= new_start_date) | (OsCard.valid_to == None))
            ).first()

            if duplicate_card:
                raise Exception(f"Kartu nomor {card_number_input} sudah aktif digunakan oleh karyawan lain.")

        target_person = OsPerson.query.get(target_emp.person_id)
        if target_person:
            target_person.name = data.get('nama', target_person.name)
            target_person.gender = data.get('gender', target_person.gender)
            target_person.pob = data.get('pob', target_person.pob)
            target_person.dob = parse_date(data.get('dob')) or target_person.dob
            target_person.religion = data.get('religion', target_person.religion)
            target_person.resident_id = data.get('resident_id', target_person.resident_id)
            target_person.address = data.get('address', target_person.address)

            if 'photo' in request.files:
                file = request.files['photo']
                if file.filename != '':
                    upload_date = datetime.now().strftime('%Y%m%d')
                    base_name = f"{employee_code_input}_{upload_date}"                    
                    saved_path = process_and_save_photo(
                        file_storage=file,
                        target_folder=UPLOAD_FOLDER,
                        filename_without_ext=base_name,
                        max_width=600,
                        quality=80
                    )
                    if saved_path:
                        target_person.photo = saved_path
            
            db.session.add(target_person)

        target_emp.employee_code = employee_code_input
        target_emp.sub_company_id = data.get('sub_company_id')
        target_emp.use_cc = use_cc_input
        target_emp.valid_from = new_start_date
        target_emp.valid_to = new_valid_to
        db.session.add(target_emp)

        # Proteksi logic jika validity child melewati validitas parent
        if c_valid_to and new_valid_to and c_valid_to > new_valid_to:
            c_valid_to = new_valid_to
        
        target_card = OsCard.query.filter_by(employee_id=id).first()
        if target_card:
            target_card.card_number = card_number_input
            target_card.valid_from = c_valid_from
            target_card.valid_to = c_valid_to
            db.session.add(target_card)
        elif card_number_input and card_number_input.lower() != 'none':
            newCard = OsCard(
                employee_id=id, 
                card_number=card_number_input, 
                valid_from=c_valid_from, 
                valid_to=c_valid_to
            )
            db.session.add(newCard)

        target_grade = OsGrade.query.filter_by(employee_id=id).first()
        if target_grade:
            target_grade.grade = data.get('grade')
            target_grade.valid_from = new_start_date
            target_grade.valid_to = new_valid_to
            db.session.add(target_grade)

        new_type_worker = data.get('type_worker')
        new_posisi = data.get('posisi')
        if (new_type_worker or new_posisi) and new_start_date:
            current_type = osType.query.filter(
                osType.employee_id == id,
                osType.valid_from <= new_start_date,
                ((osType.valid_to >= new_start_date) | (osType.valid_to == None))
            ).order_by(osType.id.desc()).first()

            if current_type:
                if current_type.type_worker != new_type_worker or current_type.posisi != new_posisi:
                    if current_type.valid_from == new_start_date:
                        current_type.type_worker = new_type_worker
                        current_type.posisi = new_posisi
                    else:
                        original_type_to = current_type.valid_to
                        current_type.valid_to = day_before
                        db.session.add(current_type)
                        
                        new_type_rec = osType(
                            employee_id=id, type_worker=new_type_worker, posisi=new_posisi,
                            valid_from=new_start_date, valid_to=original_type_to
                        )
                        db.session.add(new_type_rec)
                else:
                    current_type.valid_to = new_valid_to or current_type.valid_to
                    db.session.add(current_type)
            else:
                new_type_rec = osType(
                    employee_id=id, type_worker=new_type_worker, posisi=new_posisi,
                    valid_from=new_start_date, valid_to=new_valid_to
                )
                db.session.add(new_type_rec)

        new_cc_id = data.get('cc_id')
        if new_cc_id and new_start_date:
            master_cc = costCenter.query.get(new_cc_id)
            
            if master_cc:
                current_cc = OsCostCenter.query.filter(
                    OsCostCenter.employee_id == id,
                    OsCostCenter.valid_from <= new_start_date,
                    ((OsCostCenter.valid_to >= new_start_date) | (OsCostCenter.valid_to == None))
                ).order_by(OsCostCenter.id.desc()).first()

                cc_changed = False 
                original_cc_to = new_valid_to

                if current_cc:
                    if current_cc.org_cc_id != master_cc.id:
                        cc_changed = True
                        if current_cc.valid_from == new_start_date:
                            current_cc.cc_id = master_cc.cost_center
                            current_cc.org_cc_id = master_cc.id
                            db.session.add(current_cc)
                        else:
                            original_cc_to = current_cc.valid_to
                            current_cc.valid_to = day_before
                            db.session.add(current_cc)

                            new_cc = OsCostCenter(
                                employee_id=id, cc_id=master_cc.cost_center, org_cc_id=master_cc.id,
                                valid_from=new_start_date, valid_to=original_cc_to
                            )
                            db.session.add(new_cc)
                    else:
                        current_cc.valid_to = new_valid_to or current_cc.valid_to
                        db.session.add(current_cc)
                else:
                    new_cc = OsCostCenter(
                        employee_id=id, cc_id=master_cc.cost_center, org_cc_id=master_cc.id,
                        valid_from=new_start_date, valid_to=new_valid_to
                    )
                    db.session.add(new_cc)

                cc_def = canteen.query.join(canteenDetail, canteen.canteen_id == canteenDetail.canteen_id)\
                                      .filter(canteenDetail.org_cc_id == master_cc.id).first()
                if cc_def:
                    current_alokasi = Alokasi.query.filter(
                        Alokasi.employee_id == id,
                        Alokasi.valid_from <= new_start_date,
                        ((Alokasi.valid_to >= new_start_date) | (Alokasi.valid_to == None))
                    ).order_by(Alokasi.id.desc()).first()

                    if current_alokasi:
                        if cc_changed or current_alokasi.canteen_id != cc_def.canteen_id:
                            if current_alokasi.valid_from == new_start_date:
                                current_alokasi.canteen_id = cc_def.canteen_id
                                db.session.add(current_alokasi)
                            else:
                                current_alokasi.valid_to = day_before
                                db.session.add(current_alokasi)

                                newAlokasi = Alokasi(
                                    employee_id=id, canteen_id=cc_def.canteen_id,
                                    valid_from=new_start_date, valid_to=original_cc_to
                                )
                                db.session.add(newAlokasi)
                        else:
                            current_alokasi.valid_to = new_valid_to or current_alokasi.valid_to
                            db.session.add(current_alokasi)
                    else:
                        newAlokasi = Alokasi(
                            employee_id=id, canteen_id=cc_def.canteen_id, 
                            valid_from=new_start_date, valid_to=original_cc_to
                        )
                        db.session.add(newAlokasi)

        db.session.commit()
        return jsonify({"status": "success", "message": "Data berhasil diperbarui!"}), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({"status": "error", "message": "Terjadi kesalahan pada server: " + str(e)}), 500

@employee_bp.route('/employee/template', methods=['GET'])
def template():
    try:
        # 1. Tarik Master Data Sub Company (Terfilter Hak Akses SSO jika ada)
        allowed_subcos = get_allowed_subcompanies()
        sub_query = db.session.query(SubCompany.sub_company_name).filter(SubCompany.sub_company_name.is_not(None))
        if allowed_subcos:
            sub_query = sub_query.filter(SubCompany.sub_company_id.in_(allowed_subcos))
        sub_company_names = [r[0].strip() for r in sub_query.order_by(SubCompany.sub_company_name.asc()).all() if r[0]]

        # 2. Tarik Master Data Department / Cost Center
        dept_query = db.session.query(costCenter.org_name).filter(costCenter.org_name.is_not(None))
        department_names = [r[0].strip() for r in dept_query.order_by(costCenter.org_name.asc()).all() if r[0]]

        # 3. Inisialisasi Workbook Excel
        wb = Workbook()
        ws = wb.active
        ws.title = 'Template_Import'

        # Header Kolom Excel
        headers = [
            "Nama", "Gender", "Tempat Lahir", "Tanggal Lahir", "Agama", 
            "NIK", "Alamat", "Employee Code", "Grade", "Sub Company", 
            "Cost Center", "Type Worker", "Posisi", "Join Date", 
            "Termination Date", "Card Number", "Card Valid From", "Card Valid To"
        ]
        ws.append(headers)

        # Baris Contoh Data
        sample_subco = sub_company_names[0] if sub_company_names else "GLB"
        sample_dept = department_names[0] if department_names else "PRODUCTION"
        sample_row = [
            "Budi Contoh", "L", "Bandung", "1995-03-20", "Islam",
            "3201234567890001", "Jl. Mawar No. 12", "123456", "1", sample_subco,
            sample_dept, "DAILYWAGE", "OPERATOR", "2026-03-10",
            "", "00012.34567", "2026-03-10", ""
        ]
        ws.append(sample_row)

        # 4. Buat Hidden Sheet "Master_Data" untuk Lookup Range
        ws_master = wb.create_sheet(title='Master_Data')

        ws_master.cell(row=1, column=1, value="Sub Company")
        for idx, name in enumerate(sub_company_names, start=2):
            ws_master.cell(row=idx, column=1, value=name)

        ws_master.cell(row=1, column=2, value="Department")
        for idx, name in enumerate(department_names, start=2):
            ws_master.cell(row=idx, column=2, value=name)

        # Sembunyikan Sheet Master_Data agar Tampilan Rapi
        ws_master.sheet_state = 'hidden'

        max_subco_row = max(len(sub_company_names) + 1, 2)
        max_dept_row = max(len(department_names) + 1, 2)

        # 5. Pasang Data Validation (Dropdown)

        # A. Dropdown Gender (Kolom B)
        dv_gender = DataValidation(type="list", formula1='"L,P"', allow_blank=True)
        dv_gender.error = 'Pilih Gender L (Laki-laki) atau P (Perempuan).'
        dv_gender.errorTitle = 'Input Tidak Valid'
        ws.add_data_validation(dv_gender)
        dv_gender.add("B2:B1000")
        # B. Dropdown Agama (Kolom E)
        dv_agama = DataValidation(
            type="list", 
            formula1='"Islam,Kristen,Katolik,Hindu,Budha,Khonghucu"', 
            allow_blank=True
        )
        dv_agama.error = 'Pilih Agama dari list dropdown yang tersedia.'
        dv_agama.errorTitle = 'Agama Tidak Valid'
        ws.add_data_validation(dv_agama)
        dv_agama.add("E2:E1000")
        # C. Dropdown Sub Company (Kolom J) - Mengacu ke Sheet Master_Data
        dv_subco = DataValidation(type="list", formula1=f"Master_Data!$A$2:$A${max_subco_row}", allow_blank=True)
        dv_subco.error = 'Pilih Sub Company dari list dropdown yang tersedia.'
        dv_subco.errorTitle = 'Sub Company Tidak Valid'
        ws.add_data_validation(dv_subco)
        dv_subco.add("J2:J1000")
        # D. Dropdown Department (Kolom K) - Mengacu ke Sheet Master_Data
        dv_dept = DataValidation(type="list", formula1=f"Master_Data!$B$2:$B${max_dept_row}", allow_blank=True)
        dv_dept.error = 'Pilih Department dari list dropdown yang tersedia.'
        dv_dept.errorTitle = 'Department Tidak Valid'
        ws.add_data_validation(dv_dept)
        dv_dept.add("K2:K1000")
        # E. Dropdown Type Worker (Kolom L)
        dv_type = DataValidation(type="list", formula1='"DAILYWAGE,PIECERATE"', allow_blank=True)
        dv_type.error = 'Pilih Type Worker dari list dropdown.'
        dv_type.errorTitle = 'Type Worker Tidak Valid'
        ws.add_data_validation(dv_type)
        dv_type.add("L2:L1000")

        # 6. Stream File Excel Ke Client
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="Template_Import_Karyawan.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return jsonify({"status": "error", "message": f"Gagal membuat template: {str(e)}"}), 500

@employee_bp.route('/employee/upload', methods=['POST'])
def upload():
    file = request.files.get('file')
    if not file:
        return jsonify({"message": 'Mohon pilih file Excel terlebih dahulu.'}), 400
    
    try:
        df = pd.read_excel(file, dtype={'Card Number': str, 'Employee Code': str, 'Grade': str, 'NIK': str})

        def clean(val):
            if pd.isna(val) or val == 'nan' or val == 'NaN':
                return None
            return val

        required_columns = ['Nama', 'NIK', 'Employee Code', 'Sub Company', 'Cost Center', 'Join Date']
        missing_cols = [col for col in required_columns if col not in df.columns]
        if missing_cols:
            return jsonify({"message": f"Format Excel salah. Kolom berikut tidak ditemukan: {', '.join(missing_cols)}"}), 400

        errors = []
        notes = []
        success_count = 0

        for index, row in df.iterrows():
            line_number = index + 2
            try:
                with db.session.begin_nested():
                    nama_input = str(row['Nama']).strip() if clean(row.get('Nama')) else None
                    nik_input = clean(row.get('NIK'))
                    emp_code_input = clean(row.get('Employee Code'))
                    use_cc_input = parse_use_cc(clean(row.get('Use CC')) or 0)

                    if not nama_input or not nik_input or not emp_code_input:
                        raise ValueError("Nama, NIK, dan Employee ID tidak boleh kosong.")

                    # --- FIX: Excel Protections Date Parsing ---
                    new_start_date = extract_excel_date(row.get('Join Date'))
                    if not new_start_date:
                        raise ValueError("Tanggal 'Join Date' tidak valid atau kosong.")
                    
                    new_valid_to = extract_excel_date(row.get('Termination Date'))
                    c_valid_from = extract_excel_date(row.get('Card Valid From'))
                    c_valid_to = extract_excel_date(row.get('Card Valid To'))
                    
                    adjusted_valid_to = new_start_date - timedelta(days=1)
                    # -------------------------------------------

                    target_person = OsPerson.query.filter(OsPerson.resident_id == nik_input).first()
                    if not target_person:
                        target_person = OsPerson.query.filter(OsPerson.name.ilike(nama_input)).first()
                    
                    if target_person:
                        person_id = target_person.person_id
                        notes.append(f"Baris {line_number}: Linked data '{nama_input}' (NIK: {target_person.resident_id}).")
                        if not target_person.resident_id:
                            target_person.resident_id = nik_input
                        
                        if OsBlacklist.query.filter_by(person_id=person_id).first():
                            raise ValueError(f"Karyawan '{nama_input}' masuk daftar BLACKLIST.")

                        active_emp = OsEmployment.query.filter(
                            OsEmployment.person_id == person_id,
                            OsEmployment.valid_from <= new_start_date,
                            ((OsEmployment.valid_to >= new_start_date) | (OsEmployment.valid_to == None))
                        ).first()

                        if active_emp and active_emp.employee_code != emp_code_input:
                            raise ValueError(f"Karyawan {target_person.name} MASIH AKTIF dengan NRP lama ({active_emp.employee_code}).")
                    else:
                        newPerson = OsPerson(
                            name=nama_input, gender=clean(row.get('Gender')),
                            pob=clean(row.get('Tempat Lahir')), dob=extract_excel_date(row.get('Tanggal Lahir')),
                            religion=clean(row.get('Agama')), resident_id=nik_input,
                            address=clean(row.get('Alamat'))
                        )
                        db.session.add(newPerson)
                        db.session.flush()
                        person_id = newPerson.person_id

                    subCom_name = str(row['Sub Company']).strip() if clean(row.get('Sub Company')) else ""
                    exist_subCom = SubCompany.query.filter(SubCompany.sub_company_name.ilike(subCom_name)).first()
                    if not exist_subCom:
                        raise ValueError(f"Sub Company '{subCom_name}' tidak terdaftar.")
                    
                    cc_name = str(row['Cost Center']).strip() if clean(row.get('Cost Center')) else ""
                    exist_cc = costCenter.query.filter(costCenter.org_name.ilike(cc_name)).first()
                    if not exist_cc:
                        raise ValueError(f"Cost Center/CC '{cc_name}' tidak ditemukan.")

                    existing_active_emp = OsEmployment.query.filter(
                        OsEmployment.employee_code == emp_code_input,
                        (OsEmployment.valid_to >= new_start_date) | (OsEmployment.valid_to == None)
                    ).first()

                    if existing_active_emp and existing_active_emp.valid_from == new_start_date:
                        raise ValueError(f"Karyawan {emp_code_input} sudah memiliki data aktif di tanggal yang sama.")

                    active_emp_pk_id = None
                    if existing_active_emp:
                        existing_active_emp.valid_to = adjusted_valid_to
                        db.session.add(existing_active_emp)
                        active_emp_pk_id = existing_active_emp.id

                        if active_emp_pk_id:
                            for M in [OsCard, OsGrade, osType, OsCostCenter, Alokasi]:
                                old_recs = M.query.filter(
                                    getattr(M, "employee_id") == active_emp_pk_id,
                                    (M.valid_to >= new_start_date) | (M.valid_to == None),
                                    M.valid_from <= adjusted_valid_to
                                ).all()
                                for r in old_recs:
                                    r.valid_to = adjusted_valid_to
                                    db.session.add(r)
                                    
                    db.session.flush()

                    newEmployment = OsEmployment(
                        employee_code=emp_code_input,
                        sub_company_id=exist_subCom.sub_company_id,
                        person_id=person_id,
                        use_cc=use_cc_input,
                        valid_from=new_start_date,
                        valid_to=new_valid_to
                    )
                    db.session.add(newEmployment)
                    db.session.flush()

                    card_num = clean_str(row.get('Card Number'))
                    if card_num and card_num.lower() != 'none':
                        db.session.add(OsCard(
                            employee_id=newEmployment.id,
                            card_number=card_num,
                            valid_from=c_valid_from,
                            valid_to=c_valid_to
                        ))

                    grade_val = clean_str(row.get('Grade'))
                    if grade_val:
                        db.session.add(OsGrade(
                            employee_id=newEmployment.id,
                            grade=grade_val,
                            valid_from=new_start_date,
                            valid_to=new_valid_to
                        ))
                    
                    type_val = clean_str(row.get('Type Worker'))
                    posisi_val = clean_str(row.get('Posisi'))
                    if type_val or posisi_val:
                        db.session.add(osType(
                            employee_id=newEmployment.id,
                            type_worker=type_val,
                            posisi=posisi_val,
                            valid_from=new_start_date,
                            valid_to=new_valid_to
                        ))

                    db.session.add(OsCostCenter(
                        employee_id=newEmployment.id,
                        cc_id=exist_cc.cost_center,
                        org_cc_id=exist_cc.id,
                        valid_from=new_start_date,
                        valid_to=new_valid_to
                    ))

                    cc_def = canteen.query.join(canteenDetail, canteen.canteen_id == canteenDetail.canteen_id)\
                                          .filter(canteenDetail.org_cc_id == exist_cc.id).first()
                    if cc_def:
                        db.session.add(Alokasi(
                            employee_id=newEmployment.id,
                            canteen_id=cc_def.canteen_id,
                            valid_from=new_start_date,
                            valid_to=new_valid_to
                        ))
                
                success_count += 1

            except ValueError as ve:
                errors.append(f"Baris {line_number}: {str(ve)}")
            except Exception as e:
                errors.append(f"Baris {line_number}: {str(e)}")

        db.session.commit()

        if success_count > 0:
            status = "success" if not errors else "partial_success"
            return jsonify({
                "status": status,
                "message": f"Berhasil mengimport {success_count} data.",
                "errors": errors,
                "notes": notes
            }), 200
        else:
            return jsonify({
                "status": "error",
                "message": "Tidak ada data yang berhasil diimport.",
                "errors": errors,
                "notes": notes
            }), 400

    except Exception as e:
        db.session.rollback()
        return jsonify({"message": f"Terjadi kesalahan fatal: {str(e)}"}), 500


@employee_bp.route('/employee/export', methods=['GET'])
def export():
    try:
        search = request.args.get('search', '', type=str)
        status = request.args.get('status', 'all', type=str)
        sub_company_id = request.args.get('sub_company', '', type=str)
        department_id = request.args.get('department', '', type=str)
        
        # 1. TANGKAP TARGET DATE
        target_date_str = request.args.get('target_date', '', type=str)
        if target_date_str:
            target_date = datetime.strptime(target_date_str, '%Y-%m-%d').date()
        else:
            target_date = datetime.now().date()

        # 2. ULTIMATE POINT-IN-TIME EXPORT QUERY (NO N+1 OVERHEAD)
        query = db.session.query(
            OsEmployment,
            OsPerson,
            costCenter.org_name.label('cc_name'),
            SubCompany.sub_company_name.label('sub_con_name'),
            OsGrade.grade,
            osType.type_worker,
            osType.posisi,
            OsCard.card_number,
            OsCard.valid_from.label('card_from'),
            OsCard.valid_to.label('card_to')
        ).join(
            OsPerson, OsEmployment.person_id == OsPerson.person_id
        ).outerjoin(
            SubCompany, OsEmployment.sub_company_id == SubCompany.sub_company_id
        ).outerjoin(
            OsCostCenter,
            and_(
                OsCostCenter.employee_id == OsEmployment.id,
                OsCostCenter.valid_from <= target_date,
                or_(OsCostCenter.valid_to >= target_date, OsCostCenter.valid_to == None)
            )
        ).outerjoin(
            costCenter, costCenter.id == OsCostCenter.org_cc_id
        ).outerjoin(
            OsCard,
            and_(
                OsCard.employee_id == OsEmployment.id,
                OsCard.valid_from <= target_date,
                or_(OsCard.valid_to >= target_date, OsCard.valid_to == None)
            )
        ).outerjoin(
            OsGrade,
            and_(
                OsGrade.employee_id == OsEmployment.id,
                OsGrade.valid_from <= target_date,
                or_(OsGrade.valid_to >= target_date, OsGrade.valid_to == None)
            )
        ).outerjoin(
            osType,
            and_(
                osType.employee_id == OsEmployment.id,
                osType.valid_from <= target_date,
                or_(osType.valid_to >= target_date, osType.valid_to == None)
            )
        )

        # 3. FILTER HAK AKSES SSO
        allowed_subcos = get_allowed_subcompanies()
        if allowed_subcos:
            query = query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))
            if sub_company_id and sub_company_id not in allowed_subcos and sub_company_id not in ['TYPE_OS', 'TYPE_VENDOR']:
                return jsonify({"status": "error", "message": "Akses ditolak"}), 403

        # 4. FILTER PENCARIAN
        if search:
            query = query.filter(
                or_(
                    OsEmployment.employee_code.cast(db.String).ilike(f"%{search}%"),
                    OsPerson.name.ilike(f"%{search}%"),
                    OsCard.card_number.ilike(f"%{search}%")
                )
            )

        # 5. FILTER AKTIF / INAKTIF (POINT-IN-TIME)
        if status == 'active':
            query = query.filter(
                and_(
                    OsEmployment.valid_from.is_not(None),
                    OsEmployment.valid_from <= target_date,
                    or_(OsEmployment.valid_to >= target_date, OsEmployment.valid_to == None)
                )
            )
        elif status == 'inactive':
            query = query.filter(
                or_(
                    OsEmployment.valid_from.is_(None),
                    OsEmployment.valid_to < target_date
                )
            )

        # 6. FILTER DINAMIS (SUB COMPANY & DEPARTMENT)
        if sub_company_id == 'TYPE_OS':
            subquery_os = db.session.query(SubCompany.sub_company_id).filter(SubCompany.type_company == 'OS')
            query = query.filter(OsEmployment.sub_company_id.in_(subquery_os))
        elif sub_company_id == 'TYPE_VENDOR':
            subquery_vendor = db.session.query(SubCompany.sub_company_id).filter(SubCompany.type_company == 'Vendor')
            query = query.filter(OsEmployment.sub_company_id.in_(subquery_vendor))
        elif sub_company_id:
            query = query.filter(OsEmployment.sub_company_id == sub_company_id)

        if department_id:
            query = query.filter(OsCostCenter.org_cc_id == department_id)
        
        filtered_employees = query.all()
        
        # 7. MAPPING KE ARRAY (Langsung dari hasil JOIN, bukan lazy loading to_dict)
        data = []
        for emp, person, cc_name, sub_con_name, grade, type_worker, posisi, card_number, card_from, card_to in filtered_employees:
            data.append({
                "Employee ID": emp.employee_code or '',
                "Name": person.name or '',
                "Gender": person.gender or '',
                "Sub Company": sub_con_name or '',
                "Department": cc_name or '',
                "Card Number": card_number or '',
                "Type Worker": type_worker or '',
                "Posisi": posisi or '',
                "Join Date": emp.valid_from.strftime('%d-%b-%Y').upper() if emp.valid_from else '',
                "Termination Date": emp.valid_to.strftime('%d-%b-%Y').upper() if emp.valid_to else '',
            })

        if not data:
            return jsonify({'status': 'error', 'message': 'Data tidak ditemukan untuk diexport.'}), 400

        df = pd.DataFrame(data)
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Data_Karyawan')        
        output.seek(0)

        filename_date = target_date.strftime('%Y-%m-%d')
        return send_file(
            output, 
            as_attachment=True, 
            download_name=f"Export_Data_{filename_date}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500

@employee_bp.route('/employee/deactivate/<int:pk_id>', methods=['PUT'])
def deactivate_employee(pk_id):
    try:
        emp = OsEmployment.query.get(pk_id)
        if not emp:
            return jsonify({"status": "error", "message": "Data tidak ditemukan"}), 404

        yesterday = datetime.now().date() - timedelta(days=1)
        emp.valid_to = yesterday
        db.session.add(emp)

        target_models = [OsCard, OsGrade, OsCostCenter, osType, Alokasi]
        for Model in target_models:
            active_records = Model.query.filter(
                Model.employee_id == pk_id,
                (Model.valid_to >= yesterday) | (Model.valid_to == None)
            ).all()            
            for rec in active_records:
                rec.valid_to = yesterday
                db.session.add(rec)
        db.session.commit()
        return jsonify({
            "status": "success", 
            "message": f"Karyawan {emp.employee_code} berhasil dinonaktifkan per tanggal {yesterday.strftime('%d %b %Y')}"
        }), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({"status": "error", "message": str(e)}), 500


@employee_bp.route("/employee/stats", methods=['GET'])
def get_employee_stats():
    try:
        now = datetime.now()
        allowed_subcos = get_allowed_subcompanies()

        # 1. Base Query Employment (Terfilter Hak Akses SSO)
        base_emp = OsEmployment.query
        if allowed_subcos:
            base_emp = base_emp.filter(OsEmployment.sub_company_id.in_(allowed_subcos))

        # Total Aktif & Tidak Aktif
        total_active = base_emp.filter((OsEmployment.valid_to >= now) | (OsEmployment.valid_to == None)).count()
        total_inactive = base_emp.filter(OsEmployment.valid_to < now).count()

        # 2. Total per Cost Center (Di-JOIN langsung ke OsEmployment agar filter Subcompany presisi)
        cc_query = db.session.query(
            OsCostCenter.org_cc_id, 
            func.count(OsCostCenter.id).label('total')
        ).join(OsEmployment, OsCostCenter.employee_id == OsEmployment.id)\
         .filter((OsCostCenter.valid_to >= now) | (OsCostCenter.valid_to == None))
        
        if allowed_subcos:
            cc_query = cc_query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))
            
        stats_cc = {row.org_cc_id: row.total for row in cc_query.group_by(OsCostCenter.org_cc_id).all()}

        # 3. Total per Subcompany
        sub_query = db.session.query(
            OsEmployment.sub_company_id,
            func.count(OsEmployment.id).label('total')
        ).filter((OsEmployment.valid_to >= now) | (OsEmployment.valid_to == None))
        
        if allowed_subcos:
            sub_query = sub_query.filter(OsEmployment.sub_company_id.in_(allowed_subcos))
            
        stats_sub = {row.sub_company_id: row.total for row in sub_query.group_by(OsEmployment.sub_company_id).all()}

        # Map Ke Nama Master (Satu kali query massal tanpa loop individual query)
        all_cc = costCenter.query.all()
        
        sub_query_master = SubCompany.query.filter(SubCompany.type_company == 'OS')
        if allowed_subcos:
            sub_query_master = sub_query_master.filter(SubCompany.sub_company_id.in_(allowed_subcos))
        all_sub = sub_query_master.all()

        cc_aktif = {cc.org_name: stats_cc.get(cc.id, 0) for cc in all_cc if stats_cc.get(cc.id, 0) > 0}
        sub_aktif = {sub.sub_company_name: stats_sub.get(sub.sub_company_id, 0) for sub in all_sub}

        return jsonify({
            "status": "success",
            "data": {
                "all_total_active": total_active,
                "all_total_inactive": total_inactive,
                "active_per_cost_center": cc_aktif,
                "active_per_subCom": sub_aktif
            }
        }), 200

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500